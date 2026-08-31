"""PDF (ReportLab) and Excel (openpyxl) renderers for an inspection report."""

from __future__ import annotations

import io
from xml.sax.saxutils import escape

from app.models import InspectionReport


def esc(value: object) -> str:
    """ReportLab paragraphs are mini-XML — user data must be escaped."""
    return escape(str(value))

HEADER_FIELDS = [
    ("Company Name", "company_name"),
    ("Customer Name", "customer_name"),
    ("Supplier", "_supplier"),
    ("Part Name", "_part_name"),
    ("Part Number", "_part_number"),
    ("Drawing Number", "_drawing_number"),
    ("Drawing Revision", "drawing_revision"),
    ("Batch / Lot No.", "batch_number"),
    ("PO Number", "po_number"),
    ("Invoice Number", "invoice_number"),
    ("Inspection Date", "inspection_date"),
    ("Shift", "shift"),
    ("Machine Number", "machine_number"),
    ("Operator Name", "operator_name"),
    ("Inspector Name", "inspector_name"),
]

TABLE_COLUMNS = [
    "S.No",
    "Inspection Parameter",
    "Specification",
    "Tolerance",
    "Inspection Method",
    "Instrument",
    "Actual Value",
    "Result",
    "Defect Description",
    "Inspector Remarks",
]


def _val(report: InspectionReport, key: str) -> str:
    if key == "_supplier":
        return report.supplier.name if report.supplier else "-"
    if key == "_part_name":
        return report.part.part_name if report.part else "-"
    if key == "_part_number":
        return report.part.part_number if report.part else "-"
    if key == "_drawing_number":
        return report.drawing.drawing_number if report.drawing else "-"
    value = getattr(report, key, None)
    return "-" if value in (None, "") else str(value)


def tolerance_text(p) -> str:
    if p.upper_tolerance is None and p.lower_tolerance is None:
        return "-"
    upper = f"+{_num(p.upper_tolerance)}" if p.upper_tolerance is not None else "—"
    lower = _num(p.lower_tolerance) if p.lower_tolerance is not None else "—"
    return f"{upper} / {lower}"


def _num(value) -> str:
    if value is None:
        return "-"
    return f"{value:.4f}".rstrip("0").rstrip(".")


def _rows(report: InspectionReport) -> list[list[str]]:
    rows = []
    for r in report.results:
        p = r.parameter
        actual = r.actual_text if p.is_attribute else _num(r.actual_value)
        result = r.result
        if result == "FAIL" and r.deviation is not None:
            result = f"FAIL (dev {_num(r.deviation)})"
        rows.append(
            [
                str(r.seq),
                p.parameter,
                f"{p.specification} {p.unit or ''}".strip(),
                tolerance_text(p),
                p.inspection_method or "-",
                (p.instrument.name if p.instrument else p.instrument_text) or "-",
                actual or "-",
                result,
                r.defect_description or "-",
                r.remarks or "-",
            ]
        )
    return rows


# --------------------------------------------------------------------------- Excel
def to_xlsx(report: InspectionReport) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Inspection Report"

    bold = Font(bold=True)
    title_font = Font(bold=True, size=14)
    head_fill = PatternFill("solid", fgColor="1E3A8A")
    head_font = Font(bold=True, color="FFFFFF")
    fail_fill = PatternFill("solid", fgColor="FEE2E2")
    pass_fill = PatternFill("solid", fgColor="DCFCE7")
    thin = Side(style="thin", color="CBD5E1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(TABLE_COLUMNS))
    ws.cell(row=1, column=1, value=f"INSPECTION REPORT — {report.report_number}").font = title_font
    ws.cell(row=1, column=1).alignment = Alignment(horizontal="center")

    row = 3
    for i, (label, key) in enumerate(HEADER_FIELDS):
        col = 1 + (i % 3) * 3
        ws.cell(row=row, column=col, value=label).font = bold
        ws.cell(row=row, column=col + 1, value=_val(report, key))
        if i % 3 == 2:
            row += 1
    if len(HEADER_FIELDS) % 3:
        row += 1

    row += 1
    for c, name in enumerate(TABLE_COLUMNS, start=1):
        cell = ws.cell(row=row, column=c, value=name)
        cell.fill, cell.font, cell.border = head_fill, head_font, border
    header_row = row

    for data in _rows(report):
        row += 1
        for c, value in enumerate(data, start=1):
            cell = ws.cell(row=row, column=c, value=value)
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=c in (2, 3, 9, 10))
        if data[7].startswith("FAIL"):
            for c in range(1, len(TABLE_COLUMNS) + 1):
                ws.cell(row=row, column=c).fill = fail_fill
        elif data[7] == "PASS":
            ws.cell(row=row, column=8).fill = pass_fill

    row += 2
    ws.cell(row=row, column=1, value="INSPECTION SUMMARY").font = title_font
    summary = [
        ("Total Parameters Checked", report.total_parameters),
        ("Passed Parameters", report.passed_parameters),
        ("Failed Parameters", report.failed_parameters),
        ("Pending Parameters", report.pending_parameters),
        ("Critical Failures", report.critical_failures),
        ("Overall Inspection Result", report.overall_result),
        ("Accepted Quantity", report.accepted_quantity),
        ("Rejected Quantity", report.rejected_quantity),
        ("Rework Quantity", report.rework_quantity),
    ]
    for label, value in summary:
        row += 1
        ws.cell(row=row, column=1, value=label).font = bold
        ws.cell(row=row, column=3, value=value)

    row += 2
    for label, value in [
        ("Inspector Signature", report.inspector_signature or "-"),
        ("Signed At", str(report.inspector_signed_at or "-")),
        ("QA Approval", report.qa_signature or "-"),
        ("QA Approved At", str(report.qa_approved_at or "-")),
        ("Digital Approval Status", report.status.upper()),
    ]:
        row += 1
        ws.cell(row=row, column=1, value=label).font = bold
        ws.cell(row=row, column=3, value=value)

    widths = [6, 30, 26, 16, 22, 24, 14, 16, 26, 26]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=header_row, column=i).column_letter].width = w
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- PDF
def to_pdf(report: InspectionReport) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title=f"Inspection Report {report.report_number}",
    )
    styles = getSampleStyleSheet()
    cell = ParagraphStyle("cell", parent=styles["BodyText"], fontSize=7, leading=9)
    head = ParagraphStyle(
        "head", parent=styles["BodyText"], fontSize=7, leading=9,
        textColor=colors.white, fontName="Helvetica-Bold",
    )
    story = []

    story.append(Paragraph(f"<b>INSPECTION REPORT — {report.report_number}</b>", styles["Title"]))
    story.append(Spacer(1, 4 * mm))

    header_cells: list[list] = []
    pairs = [(label, _val(report, key)) for label, key in HEADER_FIELDS]
    for i in range(0, len(pairs), 3):
        chunk = pairs[i : i + 3]
        row: list = []
        for label, value in chunk:
            row += [Paragraph(f"<b>{esc(label)}</b>", cell), Paragraph(esc(value), cell)]
        while len(row) < 6:
            row.append("")
        header_cells.append(row)

    header_table = Table(header_cells, colWidths=[35 * mm, 55 * mm] * 3)
    header_table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f8fafc")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 2),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
            ]
        )
    )
    story += [header_table, Spacer(1, 5 * mm)]

    data: list[list] = [[Paragraph(esc(c), head) for c in TABLE_COLUMNS]]
    fail_rows: list[int] = []
    for idx, r in enumerate(_rows(report), start=1):
        data.append([Paragraph(esc(v), cell) for v in r])
        if r[7].startswith("FAIL"):
            fail_rows.append(idx)

    widths = [10, 44, 38, 24, 32, 34, 20, 26, 34, 34]
    table = Table(data, colWidths=[w * mm for w in widths], repeatRows=1)
    style = [
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e3a8a")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
    ]
    for r in fail_rows:
        style.append(("BACKGROUND", (0, r), (-1, r), colors.HexColor("#fee2e2")))
    table.setStyle(TableStyle(style))
    story += [table, Spacer(1, 6 * mm)]

    summary_rows = [
        ["Total Checked", report.total_parameters, "Passed", report.passed_parameters,
         "Failed", report.failed_parameters, "Critical Failures", report.critical_failures],
        ["Accepted Qty", report.accepted_quantity, "Rejected Qty", report.rejected_quantity,
         "Rework Qty", report.rework_quantity, "Overall Result", report.overall_result],
    ]
    summary_table = Table(
        [[Paragraph(f"<b>{esc(c)}</b>", cell) if i % 2 == 0 else Paragraph(esc(c), cell)
          for i, c in enumerate(row)] for row in summary_rows],
        colWidths=[30 * mm, 25 * mm] * 4,
    )
    summary_table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f1f5f9")),
            ]
        )
    )
    story += [Paragraph("<b>FINAL INSPECTION SUMMARY</b>", styles["Heading3"]), summary_table]
    story.append(Spacer(1, 8 * mm))

    sign_table = Table(
        [
            [
                Paragraph(
                    f"<b>Inspector</b><br/>{esc(report.inspector_signature or '—')}<br/>"
                    f"{esc(report.inspector_signed_at or '')}", cell,
                ),
                Paragraph(
                    f"<b>QA Approval</b><br/>{esc(report.qa_signature or '—')}<br/>"
                    f"{esc(report.qa_approved_at or '')}", cell,
                ),
                Paragraph(
                    f"<b>Digital Approval Status</b><br/>{esc(report.status.upper())}<br/>"
                    f"{esc(report.qa_remarks or '')}", cell,
                ),
            ]
        ],
        colWidths=[85 * mm, 85 * mm, 85 * mm],
        rowHeights=[22 * mm],
    )
    sign_table.setStyle(
        TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP")])
    )
    story.append(sign_table)

    doc.build(story)
    return buf.getvalue()
